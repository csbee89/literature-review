"""구글시트용 추출 양식(xlsx)을 만든다. 구글 드라이브에 올려 '구글 스프레드시트로 열기'로 쓴다.
시트: 안내 / 추출(지표마다 1행, 예시 2행) / 칸 설명 / 목록값.
열은 templates/extraction_form.csv의 핵심 열과 1:1로 대응한다(칸 설명 시트의 '대응 열').
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = "handover/유동성리뷰_추출시트_양식.xlsx"
F = "맑은 고딕"
HEAD = PatternFill("solid", fgColor="1F3864")
GRP = PatternFill("solid", fgColor="D9E2F3")
EX = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

# (열 이름, 묶음, 무엇을 쓰나, 없을 때, 대응 열, 너비)
COLS = [
    ("논문 번호", "기본", "배정표의 p번호", "-", "pid", 9),
    ("정리자", "기본", "본인 이름", "-", "extractor", 8),
    ("지표 이름", "지표", "논문이 쓰는 이름 그대로. 지표가 여러 개면 행을 나눈다", "-", "indicator_name", 18),
    ("정의 원문 (쪽)", "지표", "저자 문장을 따옴표 안에 그대로 + 쪽수", "보고 안 됨", "definition_verbatim", 40),
    ("단계", "분류", "S1~S5. 배정표 값이 맞으면 그대로, 다르면 고치고 비고에 이유", "-", "stage", 7),
    ("기능", "분류", "F1~F3", "-", "function", 7),
    ("시장", "분류", "MLS / 비MLS / 경매 / 복수", "-", "regime", 8),
    ("단위", "지표", "일·주·건수·비율·지수 등. 가공 뒤 단위면 그렇게 적는다", "보고 안 됨", "unit", 14),
    ("시작 사건", "지표", "기간 지표만. 예: 매물 등록일", "해당 없음", "start_event", 12),
    ("종료 사건", "지표", "기간 지표만. 예: 계약일·철회일", "해당 없음", "end_event", 12),
    ("안 팔린·철회 매물 처리", "지표", "표본에서 뺐나, 중도절단으로 뒀나, 경쟁위험으로 뒀나 (쪽)", "보고 안 됨 / 해당 없음", "censoring", 24),
    ("재등록 처리", "지표", "같은 집이 다시 올라오면 이어 붙였나, 새 매물로 셌나 (쪽)", "보고 안 됨 / 해당 없음", "relisting_handling", 22),
    ("가공·집계", "지표", "누적·로그·차분·추세 제거·계절조정·평균/중위값 등 (쪽)", "없음", "aggregation", 22),
    ("자료 이름", "자료", "예: 한국부동산원, TRREB MLS, 네이버부동산", "보고 안 됨", "data_source_name", 18),
    ("자료 기준", "자료", "신고 / 등기 / 호가(매물) / MLS / 설문 / 검색량 등", "보고 안 됨", "data_source", 14),
    ("국가·지역", "자료", "나라와 지역 단위", "-", "country; region_scale", 16),
    ("기간", "자료", "시작~끝 (월 단위까지)", "보고 안 됨", "period_start; period_end", 14),
    ("관측 수 (쪽)", "자료", "논문에 적힌 숫자 그대로. 직접 계산했으면 비고에", "보고 안 됨", "n_obs", 16),
    ("분석 방법", "방법", "모형 이름. 예: VAR·그랜저, 생존분석, 헤도닉", "-", "method_family", 20),
    ("결과 변수", "결과", "지표가 무엇과 관련되나. 예: 가격상승률", "해당 없음", "outcome_var", 16),
    ("방향", "결과", "+ / − / 없음 / 섞임. 선행이면 '지표 선행'", "해당 없음", "effect_direction", 12),
    ("크기 원문 (표·쪽)", "결과", "계수·유의수준을 표에서 그대로. 저자 문장과 표가 다르면 둘 다", "보고 안 됨", "effect_size_verbatim", 36),
    ("선행 시차", "결과", "숫자 + 단위. 예: 3개월", "해당 없음", "lead_lag", 14),
    ("검정", "결과", "그랜저, 교차스펙트럼, 표본 외 예측 등", "해당 없음", "lead_test", 14),
    ("국면·집단 차이", "결과", "상승기/하락기, 지역, 가격대별 차이 (쪽)", "보고 안 됨", "phase_dependence; heterogeneity_dims", 22),
    ("q1 정의", "질 평가", "지표 정의가 재현할 만큼 명확한가 0/1/2", "-", "q1_definition", 6),
    ("q2 선택", "질 평가", "안 팔린 매물 등 선택 편의 처리 0/1/2 (집계 지표면 NA)", "NA", "q2_selection", 6),
    ("q3 내생성", "질 평가", "역인과·누락변수 대응 0/1/2", "-", "q3_endogeneity", 6),
    ("q4 국면", "질 평가", "시장 국면을 구분했나 0/1/2", "-", "q4_phase", 6),
    ("q5 이질성", "질 평가", "지역·유형 등 나눠 봤나 0/1/2", "-", "q5_heterogeneity", 6),
    ("q6 재현성", "질 평가", "자료·코드 공개 0/1/2", "-", "q6_reproducibility", 6),
    ("비고", "기본", "분류를 고친 이유, 계산한 값, 표 오기 등", "-", "notes", 36),
]

EXAMPLES = [
    ["p115", "(예시)", "12개월 누적 아파트 거래량 (VolN)",
     '"the twelvemonth accumulated monthly apartment trading volume" (57)',
     "S4", "F3", "비MLS", "HP 필터 추세 제거 후 표준화한 값 (57, 각주 3)", "해당 없음", "해당 없음", "해당 없음", "해당 없음",
     "12개월 누적 → HP 필터 추세 제거(λ 보고 안 됨) → 표준화 (57)", "한국부동산원", "보고 안 됨 (신고·등기 미기재)",
     "한국 · 전국/수도권/5대광역시/비수도권", "2007.01~2018.10", "보고 안 됨",
     "스펙트럼 분석(교차스펙트럼) + VAR·그랜저", "아파트 가격지수 전년동월비 상승률 (57)", "지표 선행 (공통 주기에서)",
     '위상 Total 20개월 −0.3, 48개월 −0.2 (표 4, 62); 결맞음 "0.58-0.87" (61). 그랜저: 거래량→가격은 시차 2~3에서만 5% 유의, 가격→거래량은 대부분 시차에서 유의 (표 2, 60). 저자 문장: "two-way causation ... regardless of regions" (60)',
     "약 3개월(48개월 주기), 약 1.9개월(20개월 주기), 저자 환산 (61)", "교차스펙트럼, 그랜저",
     '비수도권은 선행 판단 불가 (67)', 1, "NA", 0, 0, 1, 1,
     "관측 수는 기간에서 계산하면 월 142개 × 4권역. 위상→개월 환산을 저자는 ÷π로 함(61), 통상 ÷2π면 절반 → 교수 판정. 표 4 둘째 SMetro는 NMetro로 보임 (62)"],
    ["p033", "(예시)", "Time on market (TOM)",
     '"the days between the listing date and the sale or termination date of a residential property" (580)',
     "S3", "F2", "MLS", "일 (580)", "매물 등록일 (580)", "판매일 또는 종료일 (580)",
     "안 팔린 매물을 종료 후 선택(철회·임대 전환·재등록)과 판매의 경쟁위험으로 모형화 (586~590). 기간 말 미종료 매물 처리는 보고 안 됨",
     "첫 등록만 분석, 재등록 뒤 기간은 이어 붙이지 않음 (583). 종료 매물의 약 70%가 같은 해 재등록 (596)",
     "매물 단위, 집계 없음", "TRREB MLS 매물 자료 (583)", "MLS 매물",
     "캐나다 · 토론토 광역권(GTA)", "2021 (583)", "첫 등록 매물 18,272건, 50%가 첫 등록에서 팔림 (583)",
     "중첩 로짓 + 경쟁위험 로그정규 AFT + Gumbel 코풀라 (586~590)", "TOM과 매물 결과(판매·철회·임대·재등록)",
     "호가 높을수록 TOM 길어짐 (+)",
     "ln 호가 계수 0.78 / 0.74 / 0.75 (표 2, 592); 코풀라 모수 1.16, 1.14, 1.72 (표 1, 591); 표본 외 R² 0.24 vs 독립모형 0.19 (595)",
     "해당 없음", "표본 외 예측 비교 (595)", "보고 안 됨", 2, 2, 1, 0, 1, 0,
     "코풀라 해석이 초록·591쪽('TOM을 줄인다')과 결론 598쪽('늘린다')에서 다름. 식 번호가 본문과 하나씩 어긋남"],
]

LISTS = {
    "단계": ["S1", "S2", "S3", "S4", "S5"],
    "기능": ["F1", "F2", "F3"],
    "시장": ["MLS", "비MLS", "경매", "복수"],
    "질": [0, 1, 2, "NA"],
}

wb = Workbook()

# 안내
ws = wb.active
ws.title = "안내"
guide = [
    "주택시장 유동성 지표 리뷰 · 추출 시트",
    "",
    "1. 이 파일을 구글 드라이브에 올리고 '구글 스프레드시트로 열기'로 연다. 연구실이 시트 하나를 같이 쓴다.",
    "2. '추출' 시트에 지표마다 1행을 쓴다. 한 논문에 지표가 3개면 3행이다. 회색 두 행은 예시이니 지우지 않는다.",
    "3. 칸마다 무엇을 쓰는지는 '칸 설명' 시트에 있다.",
    "4. 규칙: 논문에 있는 것만 쓴다. 논문에 없으면 「보고 안 됨」, 해당 없으면 「해당 없음」. 빈칸으로 두지 않는다.",
    "5. 숫자는 원문 그대로 옮기고 괄호 안에 쪽수를 적는다. 반올림하거나 바꿔 쓰지 않는다.",
    "6. 워드 1쪽 정리(분류·요약·한계·막힌 것)와 이 시트를 함께 낸다. 마감은 화요일 18:00.",
    "7. 단계·기능·시장 칸은 목록에서 고른다. 배정표 값과 다르게 고르면 비고에 이유를 한 줄 적는다.",
]
for i, t in enumerate(guide, 1):
    c = ws.cell(row=i, column=1, value=t)
    c.font = Font(name=F, size=14 if i == 1 else 11, bold=(i == 1))
ws.column_dimensions["A"].width = 110

# 추출
ws = wb.create_sheet("추출")
groups = []
for j, (name, grp, *_rest) in enumerate(COLS, 1):
    w = _rest[-1]
    ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = w
    g = ws.cell(row=1, column=j, value=grp)
    g.font = Font(name=F, size=9, color="1F3864", bold=True)
    g.fill = GRP
    g.border = BOX
    h = ws.cell(row=2, column=j, value=name)
    h.font = Font(name=F, size=10, color="FFFFFF", bold=True)
    h.fill = HEAD
    h.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    h.border = BOX
ws.row_dimensions[2].height = 32
for r, row in enumerate(EXAMPLES, 3):
    for j, v in enumerate(row, 1):
        c = ws.cell(row=r, column=j, value=v)
        c.font = Font(name=F, size=9)
        c.fill = EX
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.border = BOX
    ws.row_dimensions[r].height = 150
for r in range(5, 205):
    for j in range(1, len(COLS) + 1):
        c = ws.cell(row=r, column=j)
        c.font = Font(name=F, size=9)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.border = BOX
ws.freeze_panes = "D3"


def col_of(name):
    return ws.cell(row=1, column=[c[0] for c in COLS].index(name) + 1).column_letter


for key, colnames in [("단계", ["단계"]), ("기능", ["기능"]), ("시장", ["시장"]),
                      ("질", ["q1 정의", "q2 선택", "q3 내생성", "q4 국면", "q5 이질성", "q6 재현성"])]:
    li = list(LISTS).index(key) + 1
    letter = ws.cell(row=1, column=li).column_letter
    dv = DataValidation(type="list", formula1=f"=목록값!${letter}$2:${letter}${len(LISTS[key]) + 1}",
                        allow_blank=True)
    ws.add_data_validation(dv)
    for n in colnames:
        L = col_of(n)
        dv.add(f"{L}3:{L}204")

# 칸 설명
ws2 = wb.create_sheet("칸 설명")
hdr = ["열 이름", "묶음", "무엇을 쓰나", "논문에 없을 때", "대응 열 (templates/extraction_form.csv)"]
for j, t in enumerate(hdr, 1):
    c = ws2.cell(row=1, column=j, value=t)
    c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
    c.fill = HEAD
    c.border = BOX
for i, (name, grp, what, none, mapped, _w) in enumerate(COLS, 2):
    for j, v in enumerate([name, grp, what, none, mapped], 1):
        c = ws2.cell(row=i, column=j, value=v)
        c.font = Font(name=F, size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.border = BOX
for L, w in zip("ABCDE", [22, 9, 60, 20, 34]):
    ws2.column_dimensions[L].width = w
ws2.freeze_panes = "A2"

# 목록값
ws3 = wb.create_sheet("목록값")
for j, (k, vals) in enumerate(LISTS.items(), 1):
    ws3.cell(row=1, column=j, value=k).font = Font(name=F, bold=True)
    for i, v in enumerate(vals, 2):
        ws3.cell(row=i, column=j, value=v).font = Font(name=F)

wb.save(OUT)
print("written", OUT, len(COLS), "columns")
