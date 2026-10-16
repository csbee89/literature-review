#!/usr/bin/env python3
"""담당자(원석)용 엑셀 작업 파일 생성. 사용: python3 scripts/build_handover_xlsx.py
입력: papers/pdf_status.csv, students/assignments.csv, review/included_studies.csv
출력: handover/유동성리뷰_담당자파일.xlsx (생성 후 recalc 필요)"""
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

F = "Arial"
H = Font(name=F, bold=True, color="FFFFFF", size=10)
HF = PatternFill("solid", fgColor="1F3A5F")
B = Font(name=F, size=10)
BB = Font(name=F, size=10, bold=True)
T = Font(name=F, size=14, bold=True)
LINK = Font(name=F, size=10, color="1F4FD1", underline="single")
YEL = PatternFill("solid", fgColor="FFF2CC")
GRY = PatternFill("solid", fgColor="EEF1F5")
thin = Side(style="thin", color="D0D5DD")
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
WR = Alignment(wrap_text=True, vertical="top")


def head_row(ws, row, headers):
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = H
        c.fill = HF
        c.border = BOX
        c.alignment = Alignment(wrap_text=True, vertical="center")


def table(ws, headers, rows, widths, start=1, input_cols=()):
    head_row(ws, start, headers)
    for i, r in enumerate(rows, start + 1):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = B
            c.border = BOX
            c.alignment = WR
            if headers[j - 1] in input_cols:
                c.fill = YEL
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=start + 1, column=1)
    ws.auto_filter.ref = f"A{start}:{get_column_letter(len(headers))}{start + len(rows)}"


def boxed(ws, r1, r2, ncol, bold_rows=()):
    for i in range(r1, r2 + 1):
        for j in range(1, ncol + 1):
            c = ws.cell(row=i, column=j)
            c.border = BOX
            c.font = BB if i in bold_rows else B


syn = list(csv.DictReader(open("papers/pdf_status.csv", encoding="utf-8-sig")))
asg = list(csv.DictReader(open("students/assignments.csv", encoding="utf-8")))
inc = list(csv.DictReader(open("review/included_studies.csv", encoding="utf-8")))

wb = Workbook()

# 1 안내
ws = wb.active
ws.title = "안내"
lines = [
    ("유동성 지표 체계적 문헌고찰 — 담당자(이원석) 작업 파일", T),
    ("작성 2026-10-13 (v11) · 연구책임자 최승비 · 원본 데이터: 저장소 csbee89/literature-review", B),
    ("", B),
    ("원석이 하는 일 (매주)", BB),
    ("① 수요일 모임 날(W1만 10/19 월): 그 주 논문 원문을 구해 단톡방·드라이브에 올린다(원문 상태가 '받음'이 아닌 것). ② 화요일 18:00 마감 후: 워드·시트가 들어왔는지, 빈칸·쪽수 누락이 없는지만 본다. 내용이 맞는지는 판단하지 않는다. ③ 주차별배정의 노란 칸을 채운다. ④ 본인 배정 논문도 정리한다.", B),
    ("내용 대조는 Claude가 한다: Claude가 원문으로 따로 정리해 둔 것(1차 추출)과 학생 정리(2차 추출)를 맞춘 대조표를 매주 만들고, 다른 곳은 교수가 판정한다. 원석은 Claude 정리를 보지 않는다.", B),
    ("", B),
    ("시트 안내", BB),
    ("사람별일정: 8명이 몇 주차에 무엇을 하는지 한눈에 보는 표. 가장 먼저 볼 시트.", B),
    ("종합148: 이번 리뷰에서 정리할 논문 전체(교수 승인 10/11). 주차·담당·원문 받는 법이 함께 있다. p115·p033은 정리 예시라 배정하지 않는다.", B),
    ("주차별배정: 매주 누가 어떤 논문을 하는지. W1~W5는 1인 1편, W6부터 1인 2편. 매주 수요일 모임에서 배정, 다음 모임 전날 화요일 18:00 마감. W1만 10/19 배포, 첫 모임 10/28. 노란 칸(원문 배포·제출일·워드·시트·접수·반려사유)을 원석이 채운다. 맨 위 회색 '예시' 행을 따라 쓴다.", B),
    ("진행현황: 주차별배정의 노란 칸을 채우면 자동으로 계산된다(수식). 매주 수요일 보고에 이 숫자를 쓴다.", B),
    ("분류틀: 단계·기능·시장 코드의 뜻과 148편의 교차표(수식).", B),
    ("선별과정: 52,006건에서 148편이 되기까지(PRISMA). 논문 3장의 숫자.", B),
    ("전체438: 제목·초록 선별을 통과한 438편과 각 편이 종합/배경/증거지도 중 어디로 갔는지.", B),
    ("", B),
    ("노란 칸 입력 규칙", BB),
    ("원문 배포: O. 워드·시트: O(냈음) / X. 접수: O(통과) / 반려 / 미제출. 반려는 빈칸이 있거나 쪽수가 없을 때만. 반려사유는 짧게(예: 정의 칸 쪽수 없음).", B),
    ("원문 상태 '도서관 전자저널'은 학교 도서관 로그인 후 DOI 링크로 받는다. 안 되면 상호대차.", B),
    ("원문 상태 '한국 학술지'는 KCI에서 논문명 검색 → 원문보기. 안 되면 RISS·DBpia.", B),
    ("원문 상태 '받음'은 저장소 papers/ 폴더에 PDF가 있다(pid 이름).", B),
    ("", B),
    ("숫자 출처", BB),
    ("선별과정 시트의 수치는 저장소 review/prisma_flow.md(2026-10-10 검색 실행 결과)에서 옮긴 값이다.", B),
]
for i, (t, f) in enumerate(lines, 1):
    ws.cell(row=i, column=1, value=t).font = f
ws.column_dimensions["A"].width = 120

# 2 종합148
ws = wb.create_sheet("종합148")
hdr = ["pid", "주차", "담당", "연도", "제목", "학술지", "DOI", "단계", "기능", "시장", "한국", "원문 상태", "원문 받는 법", "비고"]
rows = []
for r in syn:
    rows.append([r["pid"], r["week"], r["person"], int(r["year"]), r["title"], r["venue"], r["doi_link"],
                 r["stage"], r["function"], r["regime"], r["korean"], r["pdf_status"], r["how_to_get"], r["note"]])
rows.sort(key=lambda x: (x[1] if x[1] != "예시" else "W00", x[0]))
table(ws, hdr, rows, [7, 7, 12, 7, 60, 30, 30, 26, 18, 20, 6, 22, 42, 40])
for i in range(2, len(rows) + 2):
    c = ws.cell(row=i, column=7)
    if c.value:
        c.hyperlink = c.value
        c.font = LINK
N = len(rows) + 1


# 2b 사람별일정
from collections import defaultdict
ws = wb.create_sheet("사람별일정", 1)
people = ["전우성", "손수연", "최찬미", "이원석", "이다윤", "오건", "이지은", "홍다연"]
role = {"전우성": "박사", "손수연": "석사", "최찬미": "석사", "이원석": "학부 3·담당", "이다윤": "학부 3", "오건": "학부 2", "이지은": "학부 2", "홍다연": "학부 2"}
weeks = sorted({a["week"] for a in asg})
g = defaultdict(list)
for a in asg:
    g[(a["person"], a["week"])].append(a)
ws.cell(row=1, column=1, value="사람별 주간 논문 — 수요일 모임에서 받아 다음 모임 전날 화요일 18:00까지 워드 1쪽 + 구글시트 행. W1만 10/19 배포, 첫 모임 10/28. 날짜는 배정일.").font = BB
head_row(ws, 3, ["이름", "구분"] + [f"{w} ({min(a['assign_date'] for a in asg if a['week'] == w)[5:]})" for w in weeks])
for i, p in enumerate(people, 4):
    ws.cell(row=i, column=1, value=p)
    ws.cell(row=i, column=2, value=role[p])
    for j, w in enumerate(weeks, 3):
        it = g.get((p, w), [])
        parts = [f"{x['pid']} {x['title'][:45]}" for x in it]
        c = ws.cell(row=i, column=j, value="\n".join(parts) if parts else "—")
        c.alignment = WR
boxed(ws, 4, 3 + len(people), 2 + len(weeks))
for i in range(4, 4 + len(people)):
    for j in range(3, 3 + len(weeks)):
        ws.cell(row=i, column=j).alignment = WR
    ws.row_dimensions[i].height = 75
ws.column_dimensions["A"].width = 9
ws.column_dimensions["B"].width = 11
for j in range(3, 3 + len(weeks)):
    ws.column_dimensions[get_column_letter(j)].width = 30
ws.freeze_panes = "C4"

# 3 주차별배정
ws = wb.create_sheet("주차별배정")
hdr = ["주차", "배정일", "마감", "담당", "pid", "단계", "기능", "시장", "한국", "제목", "학술지", "DOI", "원문 상태", "원문 배포", "제출일", "워드", "시트", "접수", "반려사유"]
ex = ["예시", "2026-10-19", "2026-10-27", "Claude", "p115", "S4", "F3", "비MLS 포털·중개", "Y",
      "Frequency domain analysis to detect common cycle in housing price and trading volume", "주택도시금융연구",
      "10.38100/jhuf.2021.6.2.53", "받음", "O", "2026-10-27", "O", "O", "O", ""]
rows = [ex] + [[a["week"], a["assign_date"], a["due_date"], a["person"], a["pid"], a["stage"], a["function"], a["regime"],
                a["korean"], a["title"], a["venue"], a["doi"], a["pdf_status"], None, None, None, None, None, None] for a in asg]
table(ws, hdr, rows, [7, 11, 11, 9, 7, 6, 6, 18, 6, 55, 26, 28, 18, 9, 11, 7, 7, 9, 30],
      input_cols=("원문 배포", "제출일", "워드", "시트", "접수", "반려사유"))
for j in range(1, len(hdr) + 1):
    ws.cell(row=2, column=j).fill = GRY
ws.cell(row=2, column=1).comment = Comment("쓰는 법을 보여 주는 예시 행. 진행현황 계산에서는 빠진다.", "Claude")
last = len(rows) + 1
dv1 = DataValidation(type="list", formula1='"O,반려,미제출"', allow_blank=True)
dv1.add(f"R3:R{last}")
dv2 = DataValidation(type="list", formula1='"O,X"', allow_blank=True)
for L in "NPQ":
    dv2.add(f"{L}3:{L}{last}")
ws.add_data_validation(dv1)
ws.add_data_validation(dv2)

# 4 진행현황
ws = wb.create_sheet("진행현황")
weeks = sorted({a["week"] for a in asg})
head_row(ws, 1, ["주차", "배정", "원문 배포", "제출", "통과(O)", "반려", "통과율"])
R = f"주차별배정!$A$3:$A${last}"
for i, wk in enumerate(weeks, 2):
    ws.cell(row=i, column=1, value=wk)
    ws.cell(row=i, column=2, value=f"=COUNTIF({R},A{i})")
    ws.cell(row=i, column=3, value=f'=COUNTIFS({R},A{i},주차별배정!$N$3:$N${last},"O")')
    ws.cell(row=i, column=4, value=f'=COUNTIFS({R},A{i},주차별배정!$O$3:$O${last},"<>")')
    ws.cell(row=i, column=5, value=f'=COUNTIFS({R},A{i},주차별배정!$R$3:$R${last},"O")')
    ws.cell(row=i, column=6, value=f'=COUNTIFS({R},A{i},주차별배정!$R$3:$R${last},"반려")')
    ws.cell(row=i, column=7, value=f'=IF(D{i}=0,"-",E{i}/D{i})').number_format = "0%"
t = len(weeks) + 2
ws.cell(row=t, column=1, value="합계")
for j, col in zip(range(2, 7), "BCDEF"):
    ws.cell(row=t, column=j, value=f"=SUM({col}2:{col}{t - 1})")
ws.cell(row=t, column=7, value=f'=IF(D{t}=0,"-",E{t}/D{t})').number_format = "0%"
boxed(ws, 2, t, 7, bold_rows=(t,))
ws.cell(row=t + 2, column=1, value="배정·제출은 논문 편수로 센다. 예시 p115·p033은 Claude가 정리했고 교수가 원문과 대조한다(배정 146편에서 빠짐).").font = B
for j, w in enumerate([8, 8, 8, 9, 8, 12, 9], 1):
    ws.column_dimensions[get_column_letter(j)].width = w

# 5 분류틀
ws = wb.create_sheet("분류틀")
ws.cell(row=1, column=1, value="분류 틀: 단계(주축) × 기능(절 내부) × 시장(교차)").font = T
leg = [("코드", "뜻", "대표 지표"),
       ("S1 수요·관심", "잠재 수요 신호", "검색량, 뉴스·소비 심리, 조회수"),
       ("S2 공급·호가 설정", "공급 쪽 깊이", "신규 매물·재고·흡수율, 호가 수준·호가 지수"),
       ("S3 매칭", "즉시성(시간)", "TOM/DOM, 매각확률, 호가 인하"),
       ("S4 체결·가격 발견", "비용·깊이", "거래량·회전율, sale-to-list, 호가–체결 격차"),
       ("S5 집계 상태", "복합", "tightness·heat, 유동성 지수"),
       ("F1 정의·구축", "지표를 정의하거나 만드는 논문", ""),
       ("F2 측정 쟁점", "절단·철회·재등록·집계 단위를 다루는 논문", ""),
       ("F3 선행성·관계", "지표가 가격·거래량을 앞서는지, 지표끼리의 관계", ""),
       ("시장", "MLS / 비MLS 포털·중개 / 경매 / 복수", "'(확인)'이 붙은 것은 전문을 보고 정한다")]
head_row(ws, 3, list(leg[0]))
for i, r in enumerate(leg[1:], 4):
    for j, v in enumerate(r, 1):
        ws.cell(row=i, column=j, value=v)
boxed(ws, 4, 3 + len(leg) - 1, 3)
r0 = 3 + len(leg) + 2
ws.cell(row=r0 - 1, column=1, value="교차표: 단계 × 시장, 단계 × 기능 (종합148 시트에서 자동 계산)").font = BB
stages = ["S1 수요·관심", "S2 공급·호가 설정", "S3 매칭(매물 시간·호가 조정)", "S4 체결(거래량·가격 발견)", "S5 집계 상태·지수"]
regs = ["MLS 시장", "비MLS 포털·중개", "경매 시장", "복수/불명"]
head_row(ws, r0, ["단계"] + regs + ["합계", "F1", "F2", "F3"])
for i, s in enumerate(stages, r0 + 1):
    ws.cell(row=i, column=1, value=s)
    for j in range(2, 6):
        L = get_column_letter(j)
        ws.cell(row=i, column=j, value=f'=COUNTIFS(종합148!$H$2:$H${N},$A{i},종합148!$J$2:$J${N},{L}${r0}&"*")')
    ws.cell(row=i, column=6, value=f"=SUM(B{i}:E{i})")
    for j, fc in zip((7, 8, 9), ("F1*", "F2*", "F3*")):
        ws.cell(row=i, column=j, value=f'=COUNTIFS(종합148!$H$2:$H${N},$A{i},종합148!$I$2:$I${N},"{fc}")')
tr = r0 + len(stages) + 1
ws.cell(row=tr, column=1, value="합계")
for j in range(2, 10):
    L = get_column_letter(j)
    ws.cell(row=tr, column=j, value=f"=SUM({L}{r0 + 1}:{L}{tr - 1})")
boxed(ws, r0 + 1, tr, 9, bold_rows=(tr,))
for j, w in enumerate([32, 14, 18, 12, 12, 9, 7, 7, 7], 1):
    ws.column_dimensions[get_column_letter(j)].width = w

# 6 선별과정
ws = wb.create_sheet("선별과정")
ws.cell(row=1, column=1, value="52,006건에서 148편까지 (PRISMA 2020)").font = T
pr = [("식별: Semantic Scholar 검색", 4351, "블록 ①②③ 불리언 검색, 1986년 이후"),
      ("식별: Crossref 검색", 16800, "28개 구문 × 관련도 상위 600"),
      ("식별: 핵심 학술지 16종 전수", 26581, "REE·JREFE·JHE·JUE 등 1986년 이후 모든 논문"),
      ("식별: 인용 추적", 876, "핵심 5편을 인용한 논문"),
      ("식별: OpenAlex 보완 검색", 3398, "제목·초록 불리언, API 키 사용"),
      ("식별 합계", "=SUM(B4:B8)", "위 다섯 줄의 합"),
      ("중복 제거 후", 37104, "DOI·제목으로 같은 논문 합침"),
      ("제목·초록 판독", 1400, "규칙(주택 맥락어+유동성 지표어)을 통과한 것을 한 편씩 읽음"),
      ("제목·초록 통과", 438, "주거용 매매시장, 학술지, 유동성 지표를 다룸"),
      ("→ 배경(이론 모형)", 37, "서론에서 인용만"),
      ("→ 증거지도", 253, "지표를 다른 요인의 결과변수로만 씀. 건수 그림만"),
      ("→ 종합 대상", 148, "지표 자체가 연구 대상. 교수 승인 2026-10-11"),
      ("확인: 배경+증거지도+종합", "=B13+B14+B15", '=IF(B16=B12,"438과 일치","불일치")'),
      ("출처", None, "저장소 review/prisma_flow.md (2026-10-10 실행)")]
head_row(ws, 3, ["단계", "건수", "무엇을 했나"])
for i, r in enumerate(pr, 4):
    for j, v in enumerate(r, 1):
        c = ws.cell(row=i, column=j, value=v)
        if j == 2:
            c.number_format = "#,##0"
boxed(ws, 4, 3 + len(pr), 3, bold_rows=(9, 15))
for j, w in enumerate([30, 12, 60], 1):
    ws.column_dimensions[get_column_letter(j)].width = w

# 7 전체438
ws = wb.create_sheet("전체438")
lab = {"synthesis": "종합", "background": "배경(이론)", "map": "증거지도"}
order = {"종합": 0, "배경(이론)": 1, "증거지도": 2}
rows = [[lab.get(r["scope"], r["scope"]), r.get("pid", ""), int(r["year"]) if str(r["year"]).isdigit() else r["year"],
         r["title"], r["venue"], r["doi"], r["tier"], r["status"]] for r in inc]
rows.sort(key=lambda x: (order.get(x[0], 3), x[1] or "zz", str(x[2])))
table(ws, ["구분", "pid", "연도", "제목", "학술지", "DOI", "지표층", "상태"], rows, [11, 7, 7, 70, 32, 32, 8, 6])

wb.save("handover/유동성리뷰_담당자파일.xlsx")
print("saved")
